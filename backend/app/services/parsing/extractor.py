"""
Document text extraction service for PDF and DOCX files.

Implements:
- pdfplumber extraction with automatic PyMuPDF fallback
- python-docx extraction (paragraphs and table cells)
- Blank or image-only scanned PDF detection (raises ScannedOrEmptyError)
"""

from __future__ import annotations

from pathlib import Path

import docx
import pdfplumber
import pymupdf

from app.core.logging import get_logger, log_event

logger = get_logger(__name__)

# Minimum characters expected from a legitimate machine-readable resume
MIN_TEXT_CHARS = 50


class ParsingError(Exception):
    """Base exception for document parsing failures."""

    def __init__(self, error_code: str, message: str) -> None:
        self.error_code = error_code
        self.message = message
        super().__init__(message)


class ScannedOrEmptyError(ParsingError):
    """Raised when document contains no extractable text (e.g. image-only PDF)."""


def extract_pdf_text(file_path: Path) -> str:
    """
    Extract text from a PDF file using pdfplumber, falling back to PyMuPDF.

    Raises ScannedOrEmptyError if total extracted text is under MIN_TEXT_CHARS.
    """
    text_chunks: list[str] = []

    # 1. Primary extractor: pdfplumber
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_chunks.append(page_text)
    except (OSError, ValueError, RuntimeError) as exc:
        log_event(logger, "pdfplumber_extraction_failed")
        logger.debug(
            "pdfplumber failed on %s, trying PyMuPDF fallback: %s", file_path.name, exc
        )

    combined_text = "\n".join(text_chunks).strip()

    # 2. Fallback to PyMuPDF if pdfplumber returned empty or insufficient text
    if len(combined_text) < MIN_TEXT_CHARS:
        try:
            doc = pymupdf.open(file_path)
            fallback_chunks: list[str] = []
            for page in doc:
                t = page.get_text()
                if t:
                    fallback_chunks.append(t)
            doc.close()
            fallback_text = "\n".join(fallback_chunks).strip()
            if len(fallback_text) > len(combined_text):
                combined_text = fallback_text
        except (OSError, ValueError, RuntimeError) as exc:
            log_event(logger, "pymupdf_extraction_failed")
            logger.debug("PyMuPDF fallback failed on %s: %s", file_path.name, exc)

    if len(combined_text) < MIN_TEXT_CHARS:
        raise ScannedOrEmptyError(
            "scanned_or_empty",
            "PDF document contains insufficient machine-readable text (likely scanned or empty).",
        )

    return combined_text


def extract_docx_text(file_path: Path) -> str:
    """
    Extract text from a DOCX document including all body paragraphs and tables.

    Raises ScannedOrEmptyError if total extracted text is under MIN_TEXT_CHARS.
    """
    try:
        doc = docx.Document(file_path)
        chunks: list[str] = []

        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                chunks.append(para.text.strip())

        # Extract table contents (tables frequently contain resume layout/sections)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if row_cells:
                    chunks.append(" | ".join(row_cells))

        combined_text = "\n".join(chunks).strip()
    except Exception as exc:
        log_event(logger, "docx_extraction_failed")
        raise ParsingError(
            "corrupted_file", "Failed to open or parse DOCX document."
        ) from exc

    if len(combined_text) < MIN_TEXT_CHARS:
        raise ScannedOrEmptyError(
            "scanned_or_empty",
            "DOCX document contains insufficient text content.",
        )

    return combined_text


def extract_document_text(file_path: Path, extension: str) -> str:
    """Route document to appropriate extractor based on extension."""
    norm_ext = extension.lower().strip()
    if not norm_ext.startswith("."):
        norm_ext = f".{norm_ext}"

    if norm_ext == ".pdf":
        return extract_pdf_text(file_path)
    elif norm_ext == ".docx":
        return extract_docx_text(file_path)
    else:
        raise ParsingError(
            "unsupported_extension", f"Unsupported document format '{norm_ext}'."
        )
